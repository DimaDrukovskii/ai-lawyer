from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile

import yaml
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from ..schemas import ClientProfile

TEXT_SUFFIXES = {".txt", ".md", ".xml", ".json", ".csv", ".tsv"}
MAX_DOC_CHARS = 60_000


class DocReadError(RuntimeError):
    pass


def _read_xlsx(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        lines: list[str] = []
        for ws in wb.worksheets:
            lines.append(f"## лист: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None and str(c).strip()]
                if cells:
                    lines.append("\t".join(cells))
        return "\n".join(lines)
    finally:
        wb.close()


def _read_csv(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="replace")
    dialect = csv.Sniffer().sniff(raw[:2048], delimiters=",;\t") if raw.strip() else csv.excel
    return "\n".join("\t".join(r) for r in csv.reader(io.StringIO(raw), dialect))


def read_doc(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            return "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)
        if suffix in {".xlsx", ".xlsm"}:
            return _read_xlsx(path)
        if suffix == ".csv":
            return _read_csv(path)
        if suffix in TEXT_SUFFIXES:
            return path.read_text(encoding="utf-8", errors="replace")
    except (
        OSError,
        csv.Error,
        ValueError,
        KeyError,
        PyPdfError,
        BadZipFile,
        InvalidFileException,
    ) as exc:
        raise DocReadError(f"{path.name}: {exc}") from exc
    raise DocReadError(f"{path.name}: неподдерживаемый формат {suffix or '(без расширения)'}")


@dataclass(frozen=True)
class Case:
    name: str
    dir: Path
    profile: ClientProfile
    context: str
    docs: dict[str, str]
    unreadable: tuple[str, ...] = ()

    @property
    def local_context_path(self) -> Path:
        return self.dir / "context.local.md"


def load_case(case_dir: Path) -> Case:
    profile_path = case_dir / "profile.yaml"
    if not profile_path.exists():
        raise FileNotFoundError(f"нет {profile_path}: опиши клиента (см. cases/demo/profile.yaml)")
    profile = ClientProfile.model_validate(yaml.safe_load(profile_path.read_text(encoding="utf-8")))

    context_parts = [
        p.read_text(encoding="utf-8")
        for p in (case_dir / "context.md", case_dir / "context.local.md")
        if p.exists()
    ]
    docs: dict[str, str] = {}
    unreadable: list[str] = []
    docs_dir = case_dir / "docs"
    for path in sorted(docs_dir.glob("*")) if docs_dir.exists() else []:
        if not path.is_file() or path.name == "README.txt":
            continue
        try:
            docs[path.name] = read_doc(path)
        except DocReadError as exc:
            unreadable.append(str(exc))
    return Case(
        case_dir.name, case_dir, profile, "\n\n".join(context_parts), docs, tuple(unreadable)
    )


def _neutralize(text: str) -> str:
    """Текст документа не должен уметь закрыть наш тег и «выйти» в служебную часть промпта."""
    return text.replace("</", "<\u200b/")


def pack(case: Case) -> str:
    """Один текстовый блок для промптов: профиль, контекст, документы."""
    docs = "\n\n".join(
        f"### {name}\n{_neutralize(text[:MAX_DOC_CHARS])}" for name, text in case.docs.items()
    )
    return (
        f"<profile>\n{case.profile.model_dump_json(indent=1)}\n</profile>\n"
        f"<context>\n{_neutralize(case.context)}\n</context>\n"
        f"<documents>\n{docs}\n</documents>"
    )
